"""Rebuild the four Handover Package documents.

Each document is rewritten in place on top of its own previous version, which
supplies the page setup, styles and footer; the helpers below reproduce the
original formatting (accent headings, navy table headers, red TEAM TO COMPLETE
boxes). Edit the content functions and re-run to produce the next version.

    python build_handover.py <handover dir>     (needs python-docx)

The previous version is archived by hand under handover/archive/vX.Y/ first.
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

VERSION = "Version 2.2 · 15 October 2026 (living document -- updated each iteration) | Team TE38 · BitDance | FIT5120, Monash University"
ACCENT, INK, MUTED, BORDER = "1F3D7A", "201E19", "6B6455", "D9D3C4"
FLAG_INK, FLAG_FILL, FLAG_BORDER = "9C2F2F", "FBEFEF", "E3B8B8"
IMG = Path(__file__).parent / "img"  # screenshots of the live site and the diagrams
LIVE = "https://handmirror.vercel.app"
REPO = "https://github.com/randlyoyo/FIT5120-TE38-SignLanguage"
API = "https://fit5120-te38-signlanguage-production.up.railway.app"
GOVERNANCE = "https://drive.google.com/drive/folders/1u6BIRZyPxSSnvfm3VgOKKIQmPg2BWyyp"
MENTORS = "Afsana and Ted"
TEAM = [  # name, role, email, LinkedIn -- from the team's Team Info document (iteration 3)
    ["Dingao Wang", "Lead developer; support contact", "dwan0104@student.monash.edu", "www.linkedin.com/in/dingaowang"],
    ["Ziyang Wang", "Team member", "zwan0637@student.monash.edu", "www.linkedin.com/in/ziyang-wang-2b4700197"],
    ["Minshu Ye", "Team member", "myee0023@student.monash.edu", "www.linkedin.com/in/minshu-ye-24887b428"],
    ["Xinyue Mao", "Team member", "xmao0020@student.monash.edu", "www.linkedin.com/in/xinyue-mao-236877428"],
    ["Yan Zhuang", "Team member", "yzhu0233@student.monash.edu", "www.linkedin.com/in/yan-zhuang-b4b80b3a3"],
    ["Ruizi Song", "Team member; holds the Sign Chat model weights", "rson0034@student.monash.edu", "www.linkedin.com/in/ruizi-song-12a8a8392"],
]


class Doc:
    def __init__(self, path: Path):
        self.path = path
        self.fig = 0
        self.d = Document(path)
        body = self.d.element.body
        for el in list(body.iterchildren()):
            if el.tag != qn("w:sectPr"):
                body.remove(el)

    # --- paragraphs ---
    def _p(self, before=None, after=None, style=None):
        p = self.d.add_paragraph(style=style)
        pf = p.paragraph_format
        if before is not None:
            pf.space_before = Pt(before / 20)
        if after is not None:
            pf.space_after = Pt(after / 20)
        return p

    @staticmethod
    def _run(p, text, color=INK, size=21, bold=False, italic=None, font=None):
        lines = text.split("\n")
        r = p.add_run(lines[0])
        for line in lines[1:]:
            r.add_break()
            r.add_text(line)
        r.font.color.rgb = RGBColor.from_string(color)
        r.font.size = Pt(size / 2)
        r.bold = bold or None
        if italic is not None:
            r.italic = italic
        if font:
            r.font.name = font
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        return r

    def header(self, kicker, title="HandMirror"):
        self._run(self._p(after=40), kicker, ACCENT, 20, bold=True)
        self._run(self._p(after=40), title, INK, 48, bold=True)
        self._run(self._p(after=280), VERSION, MUTED, 20, italic=True)

    def h1(self, text):
        p = self._p(before=280, after=120)
        p.paragraph_format.keep_with_next = True
        self._run(p, text, ACCENT, 28, bold=True)

    def h2(self, text):
        p = self._p(before=160, after=80)
        p.paragraph_format.keep_with_next = True
        self._run(p, text, INK, 23, bold=True)

    def para(self, text):
        self._run(self._p(after=160), text, INK, 21, italic=False)

    def code(self, text):
        self._run(self._p(after=160), text, INK, 18, italic=False, font="Consolas")

    def bullets(self, items):
        for it in items:
            p = self._p(after=60, style="List Bullet")
            self._run(p, it, INK, 21)

    def figure(self, name, caption, width=16.0):
        """A centred image from img/ with a numbered caption under it."""
        self.fig += 1
        p = self._p(before=80, after=40)
        p.alignment = 1
        p.paragraph_format.keep_with_next = True  # the caption stays under its picture
        p.add_run().add_picture(str(IMG / name), width=Cm(width))
        cap = self._p(after=200)
        cap.alignment = 1
        self._run(cap, f"Figure {self.fig}. {caption}", MUTED, 18, italic=True)

    def closing(self, text):
        self._run(self._p(before=320), text, MUTED, 18)

    # --- tables ---
    @staticmethod
    def _borders(tbl, color):
        tblPr = tbl._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
            b.append(e)
        tblPr.append(b)

    @staticmethod
    def _cell(cell, width, fill, margins):
        tcPr = cell._tc.get_or_add_tcPr()
        w = OxmlElement("w:tcW")
        w.set(qn("w:type"), "dxa")
        w.set(qn("w:w"), str(width))
        tcPr.append(w)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)
        mar = OxmlElement("w:tcMar")
        for side, v in zip(("top", "bottom", "start", "end"), margins):
            e = OxmlElement(f"w:{side}")
            e.set(qn("w:w"), str(v))
            e.set(qn("w:type"), "dxa")
            mar.append(e)
        tcPr.append(mar)

    def table(self, headers, rows, widths):
        t = self.d.add_table(rows=1 + len(rows), cols=len(headers))
        t.alignment = 1  # centre
        self._borders(t, BORDER)
        for r, values in enumerate([headers] + rows):
            for c, value in enumerate(values):
                cell = t.cell(r, c)
                self._cell(cell, widths[c], ACCENT if r == 0 else "FFFFFF", (60, 60, 120, 120))
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)  # the style's spacing made every row a line taller
                self._run(p, value, "FFFFFF" if r == 0 else INK, 19, bold=(r == 0))
        self._p(after=80)

    def flag(self, text):
        t = self.d.add_table(rows=1, cols=1)
        self._borders(t, FLAG_BORDER)
        cell = t.cell(0, 0)
        self._cell(cell, 9972, FLAG_FILL, (80, 80, 140, 140))
        p = cell.paragraphs[0]
        self._run(p, "TEAM TO COMPLETE: ", FLAG_INK, 19, bold=True)
        self._run(p, text, FLAG_INK, 19)
        self._p(after=80)

    def save(self):
        self.d.save(self.path)


# ---------------------------------------------------------------- content ---

SHARED_LICENCE_ROWS = [
    ["Sign records & demonstration videos", "MM-WLAuslan word-level dataset (UQ CV Lab); dictionary cross-reference to Auslan Signbank", "Research dataset terms; Signbank is CC BY-NC-ND 4.0"],
    ["Practice recognition model", "Trained by the team on MM-WLAuslan", "Inherits the dataset's terms"],
    ["3D avatar body", "SMPL-X neutral model (Max Planck Institute for Intelligent Systems)", "Non-commercial use only; the model may not be redistributed"],
    ["Avatar motion", "SMPL-X parameters fitted by the team to the MM-WLAuslan signers", "Inherits the dataset's terms"],
    ["Avatar skin", "BEDLAM body textures by Meshcapade", "CC BY-NC 4.0"],
    ["Avatar clothing", "Meshcapade SMPL texture samples", "CC BY 4.0"],
    ["Avatar hair", "Exported from an Avaturn avatar", "CC BY-NC 4.0 -- non-commercial use only"],
    ["Scenario pictograms", "Human Pictogram 2.0 by TopeconHeroes (recoloured)", "Free to use as part of a design; files may not be redistributed or sold on their own"],
    ["Sign Chat models", "Uni-Sign (sign-to-text) and SignSparK (text-to-sign), fine-tuned on Auslan-Daily; Qwen2.5-1.5B-Instruct for dialogue", "Per each upstream project's licence"],
    ["Home page cover footage", "Four clips of people signing from Pexels, video IDs 10373883, 10374296, 10374298, 10373841 (client/public/hero)", "Pexels licence: free to use, attribution not required"],
]


def product(d: Doc):
    d.header("HANDOVER PACKAGE — PRODUCT DOCUMENT")
    d.h1("1. Overview")
    d.para(
        "HandMirror is a free, web-based Auslan (Australian Sign Language) learning assistant for families. "
        "This version delivers the full learning loop the team planned: look up a sign; learn it three ways "
        "(a real signer on video, a step-by-step breakdown, and a 3D avatar you can turn to any angle); practise "
        "it on camera with instant feedback; organise learning with personalised To Learn lists and real-life "
        "scenario journeys; and hold a short conversation with a signing avatar (Sign Chat)."
    )
    d.figure("01-home.jpg", "The home page states what HandMirror is and who it is for within the first screen.", 10.5)
    d.h2("The problem")
    d.para(
        "Parents and caregivers of young Deaf or hard-of-hearing children in Australia often have limited Auslan "
        "knowledge and struggle to find practical, trustworthy learning support for everyday family communication. "
        "Australian news coverage (ABC News, 2022) documented a family that secured NDIS funding but still could not "
        "find a suitable Auslan tutor -- funding does not always create access. Existing online resources are "
        "scattered, often move too fast for a beginner, and cannot confirm whether a learner is practising a sign "
        "correctly."
    )
    d.h2("The audience")
    d.para(
        "Primary audience: parents and caregivers of young Deaf or hard-of-hearing children in Australia, with "
        "limited Auslan knowledge and a need for flexible, family-friendly support. Primary persona: “Sarah "
        "Thompson”, 38, a parent and primary-school teacher who needs short, flexible learning sessions with "
        "progress tracking. Secondary audiences: Auslan beginners generally, and accessibility-focused learners "
        "(students, teachers)."
    )
    d.h2("Why it's important")
    d.para(
        "The project addresses UN Sustainable Development Goal 4 (Quality Education) and Goal 10 (Reduced "
        "Inequalities) -- giving families a flexible, low-cost way to build practical Auslan at home, rather than "
        "depending on scarce tutor availability."
    )

    d.h1("2. What's built so far")
    d.table(
        ["Epic", "Status", "What it does today"],
        [
            ["Epic 1 -- Sign library search & browse", "Delivered",
             "Keyword search across gloss, tags, keywords and definitions; topic browsing; difficulty-level filter and sort; pagination; "
             "sign-to-search (sign at the webcam and pick from the four closest signs); select several signs at once and add them to "
             "To Learn or mark them learnt."],
            ["Epic 2 -- Sign demonstration & learning", "Delivered",
             "Real signer video for every sign (0.5×–2× speed), a step-by-step movement description, a rotatable 3D avatar "
             "playing the sign, Mark as learnt, and + / ✓ buttons on every sign card for one-click saving."],
            ["Epic 3 -- Basic quiz / self-check", "Not built", "Remains open for a future iteration."],
            ["Epic 4 -- AI sign recognition & webcam feedback", "Delivered",
             "Practice tab: the learner signs at the webcam and gets one of three results -- Excellent (score 80+), "
             "Almost there (60–79) or Keep practicing (under 60) -- with a tip. Either hand may lead."],
            ["Epic 5 -- Personalised learning", "Delivered",
             "To Learn list (grouped by when signs were added), a personalise wizard (pick topics, a word count and its split; "
             "it fills the list with signs not yet learnt), and a Learnt list with progress."],
            ["Epic 6 -- Family mode", "Not built", "Remains open for a future iteration."],
            ["Epic 7 -- Real-life conversation practice", "Delivered",
             "Scenario Learning: 34 journeys (e.g. café, doctor, family) split into short stops, with a guided route and progress. "
             "Sign Chat: type or sign a sentence and a 3D avatar replies in Auslan."],
        ],
        [2600, 1300, 6072],
    )

    d.h1("3. How it works (high level)")
    d.para(
        "A React + TypeScript client (hosted on Vercel) calls a Node.js + Express REST API (hosted on Railway), backed "
        "by a MySQL database (also on Railway). Demonstration videos are served from Cloudflare R2. Practice runs the "
        "hand-tracking in the learner's browser and scores the attempt with a small recognition model inside the "
        "Express API. Sign Chat uses a separate Python model service that needs a GPU. Full technical detail is in "
        "the Support and Maintenance Documents."
    )
    d.para(f"Live product: {LIVE} (public -- no account, login or password is needed)   |   Source code: {REPO}")

    d.h1("4. Ethics & responsible design")
    d.h2("Data provenance and licensing")
    d.para(
        "Every sign record carries a source field recording where its content came from, so the product never "
        "silently presents unverified content as confirmed Auslan. The third-party assets in the product, and the "
        "terms they come under, are:"
    )
    d.table(["Asset", "Source", "Licence / terms"], SHARED_LICENCE_ROWS, [2400, 4300, 3272])
    d.para(
        "Licence position: the restricted assets above -- Auslan Signbank, the research datasets and everything derived "
        "from them, SMPL-X, the BEDLAM skin and the avatar hair -- are licensed under Creative Commons 4.0 "
        "non-commercial terms (CC BY-NC 4.0; CC BY-NC-ND 4.0 for Signbank) or an equivalent non-commercial research "
        "licence. HandMirror may therefore be run free of charge, as it is today, for non-commercial use only. Any "
        "commercial use requires a separate licence from each rights holder first."
    )
    d.h2("Content safety")
    d.para(
        "The dictionary covers the full Auslan vocabulary of its source dataset, which includes adult and crude signs -- "
        "appropriate for a reference, but not for a family homepage. The home page's sign strip therefore draws only "
        "from a fixed list of 16 family-friendly signs (HELLO, THANK YOU, DOG, MOTHER and so on), never from a random "
        "catalogue pick. The cover footage behind the home page title is stock video of people signing, chosen for "
        "mood; it is not presented as a lesson."
    )
    d.h2("Privacy")
    d.bullets([
        "No user accounts. Learnt and To Learn lists are stored only in the visitor's own browser (localStorage) and are never sent to the server.",
        "Practice: camera video never leaves the browser. Hand-tracking runs locally; only numeric landmark coordinates are posted for scoring, and they are not stored.",
        "Sign Chat: typed messages are sent to the Sign Chat service. When a learner chooses to sign a message instead, the recorded clip is uploaded to that service so it can be recognised.",
    ])
    d.para("What the Sign Chat service keeps (checked in sign_chat_backend/signchat, 13 October 2026):")
    d.bullets([
        "Uploaded video clips: written to a temporary file for recognition and deleted as soon as the request finishes, whether it succeeds or fails.",
        "Conversation text: the last 6 exchanges of each session, in memory only (at most 1,000 sessions). Lost when the service restarts; DELETE /api/session/:id clears one sooner.",
        "Generated replies: each reply's motion file, skeleton video and subtitles (which contain the reply text) are saved under media/ and are not cleaned up automatically. On Colab they disappear when the runtime ends; a permanent host should delete them after a set period.",
        "No accounts, names or IP addresses are recorded by the service itself.",
    ])
    d.para(
        "Personal information is handled in line with the Australian Privacy Principles (Privacy Act 1988, Cth): "
        "nothing is collected that the feature does not need, camera video is not stored, and no learner is identified. "
        "The mapping to each principle is in the Support Document, Section 7. Recommendation: tell users on the Sign "
        "Chat page that their messages are processed by the service and not kept beyond the session."
    )
    d.h2("Honesty about limits")
    d.para(
        "The interface tells learners that the 3D avatar shows the overall movement and that exact handshapes should be "
        "taken from the demonstration video. The Practice score threshold was calibrated on studio footage, and Sign "
        "Chat's replies have not yet been checked by fluent signers -- see the Support Document's known limitations."
    )
    d.h2("ACS Code of Conduct mapping")
    d.table(
        ["ACS value", "How HandMirror addresses it"],
        [
            ["The public interest", "Targets a documented, real accessibility gap for Deaf/hard-of-hearing families rather than a purely commercial idea."],
            ["Quality of life", "Free to use, no account required, reduces dependence on scarce and costly in-person tutoring."],
            ["Honesty", "Source fields disclose provenance; the avatar is labelled as a guide; adult signs are kept off the family-facing home page."],
            ["Competence", "Recognition has a parity test against the training pipeline; the APIs and data paths are documented for safe extension."],
            ["Professional development", "The recognition contract (recognition/API.md) and Sign Chat README let future contributors retrain or replace models without redesigning the product."],
            ["Professionalism", "Known limitations and licence restrictions (non-commercial use only) are disclosed rather than hidden."],
        ],
        [2400, 7572],
    )
    d.para("The full Ethics Canvas is in the Project Governance Portfolio: " + GOVERNANCE)

    d.h1("5. Team")
    d.table(["Name", "Role", "Email", "LinkedIn"], TEAM, [2000, 2400, 2700, 2872])
    d.para("Team TE38 · “BitDance” · FIT5120, Monash University · Studio mentors: " + MENTORS + ".")

    d.h1("6. Potential sponsors")
    d.table(
        ["Organisation", "Why they fit"],
        [
            ["Deaf Children Australia", "Runs family programs and Auslan resources; could validate family needs and test the platform directly with parents and caregivers."],
            ["Expression Australia", "Provides Auslan education and services (incl. Auslan in the Home, Auslan Storytime); could advise on Auslan accuracy, accessibility and culturally appropriate design."],
        ],
        [2600, 7372],
    )

    d.h1("7. Links")
    d.table(
        ["Item", "Link"],
        [
            ["Live product", LIVE + " -- public, no login or password"],
            ["Source code (GitHub)", REPO],
            ["System architecture", "Current diagram: Support Document, Section 3 (and Maintenance Document, Section 2); earlier write-up in website/docs/system-architecture.docx"],
            ["Product demo video", "Project Governance Portfolio → Project Support folder"],
            ["Team LinkedIn", "Each member's profile is listed in Section 5"],
            ["Project Governance Portfolio (Google Drive)", GOVERNANCE],
        ],
        [3400, 6572],
    )
    d.closing("HandMirror — Product Document · Team TE38 (BitDance) · FIT5120")


def support(d: Doc):
    d.header("HANDOVER PACKAGE — SUPPORT DOCUMENT")
    d.h1("1. Purpose & audience")
    d.para(
        "This document is for whoever operationally runs and supports the live HandMirror product once a sponsor has "
        "taken it on: what every feature does (so support staff can help learners), day-to-day operation, data "
        "loading, backups, security, and what to do when something breaks. It assumes no prior familiarity with the "
        "codebase -- that is the companion Maintenance Document."
    )
    d.para(
        "Reading order: Section 2 for the features, Sections 3-6 to run the system, Section 7 for security and the "
        "standards it follows, Section 9 when something goes wrong, and Section 10 for the handover plan itself."
    )

    d.h1("2. Using HandMirror -- feature guide")
    d.para(
        "HandMirror needs no account: every feature works as soon as the site opens. A learner's To Learn and Learnt "
        "lists are kept in their own browser, so they are tied to that browser and device."
    )
    d.h2("2.1 Look up a sign (Library)")
    d.bullets([
        "Type any English word in the search bar; results match the sign's name, its topics, search keywords and definitions.",
        "Narrow the list with the topic index on the left, the All levels / Beginner / Intermediate / Advanced switch, and Sort by (A-Z, difficulty or popularity).",
        "The camera icon in the search bar is sign-to-search: sign at the webcam and pick from the four closest matches.",
        "Each card plays its demonstration video, shows the estimated level (top-left), and has two buttons at the bottom-right: + adds the sign to To Learn, ✓ marks it learnt. A blue “To Learn” or green “Learnt” stamp shows its state.",
    ])
    d.figure("02-library.jpg", "The Sign Library: search, topic index, level filter and sort, with the + / ✓ buttons on each card.")
    d.h2("2.2 Save several signs at once")
    d.para(
        "Select signs (above the cards in the Library, To Learn and Learnt) turns the grid into a multi-select: click "
        "cards to tick them, then use the bar at the bottom of the screen. The Library offers Add to To Learn and Mark "
        "as learnt; To Learn offers Mark as learnt and Remove; Learnt offers Move to To Learn and Remove. A sign is never "
        "on both lists -- adding it to one takes it off the other. Done leaves selection mode."
    )
    d.figure("03-library-select.jpg", "Two signs selected, with the bulk-action bar at the bottom of the screen.")
    d.h2("2.3 Learn a sign")
    d.bullets([
        "The sign page opens on Learn: the demonstration video (with a choice of versions where there are several, replay, and 0.5×-2× speed), the definition, and “How it's formed” -- the movement step by step.",
        "Video / 3D avatar switches the stage to an avatar performing the same sign. Drag to turn it to any angle; Reset view puts it back. The avatar shows the overall movement -- exact handshapes come from the video.",
        "Add to to-learn list and Mark as learnt save the sign; the arrows at the sides step to the previous and next sign in the list the learner came from.",
    ])
    d.figure("04-sign-learn.jpg", "A sign page in Learn mode with the demonstration video.", 15)
    d.figure("05-sign-avatar.jpg", "The same sign on the 3D avatar.", 15)
    d.h2("2.4 Practise a sign")
    d.para(
        "Practice (top-right of the sign page) is shown for the signs the recognition model knows. Read the note above "
        "Start practice: follow the hand orientation in the video. After a short countdown the learner signs at the "
        "webcam; the capture stops by itself when the hands come to rest. The result is Excellent (80+), Almost there "
        "(60-79) or Keep practicing (under 60), with a tip and Try again. The camera picture never leaves the browser "
        "(Section 7)."
    )
    d.figure("06-practice.jpg", "Practice mode, before the camera starts.", 15)
    d.h2("2.5 To Learn and the personalise wizard")
    d.para(
        "To Learn lists the saved signs grouped by when they were added (Day / Week / Month / Year). Personalize your "
        "learning opens a two-step wizard: pick topics, then set how many signs in total and how they split across "
        "the topics (drag the pie chart or type numbers). The wizard adds signs the learner has not saved or learnt yet."
    )
    d.figure("07-to-learn.jpg", "The To Learn list, with the personalise card above it.", 15)
    d.figure("08-personalise.jpg", "Step 1 of the personalise wizard: choosing topics.", 15)
    d.h2("2.6 Learnt")
    d.para("Learnt shows every sign marked learnt, grouped the same way, and progress against the whole dictionary.")
    d.figure("09-learnt.jpg", "The Learnt list and its progress bar.", 15)
    d.h2("2.7 Scenario journeys")
    d.para(
        "Scenarios groups signs by real situations -- ordering a coffee, a doctor's visit, family -- in three "
        "directions (Daily Life, Professional, Language Basics) that open and close with the arrow on the left. "
        "Each journey is a route of short stops. Hovering over a word shows its video with three buttons: Practise "
        "this sign, Add to To Learn and Mark as learnt. Next stop moves along the route."
    )
    d.figure("10-scenarios.jpg", "The Scenarios page: how a journey works, and the first direction opened.", 15)
    d.figure("11-journey.jpg", "One journey: the route of stops on the left, the current stop's signs on the right.", 15)
    d.h2("2.8 Sign Chat")
    d.para(
        "Type a short English sentence (or press the record button to sign it on camera) and the avatar replies in "
        "Auslan with subtitles; click a reply to watch it again, and New chat to start over. Sign Chat depends on the "
        "GPU model service in Section 4 -- when that is not running, the page shows an error instead of a reply."
    )
    d.figure("12-sign-chat.jpg", "Sign Chat waiting for a message.", 15)

    d.h1("3. System summary")
    d.figure("arch.png", "How the parts connect. The browser talks to four services; only the Sign Chat service is not permanently hosted.")
    d.table(
        ["Component", "Technology", "Where it runs"],
        [
            ["Web client", "React 19 + TypeScript (Vite); three.js for the 3D avatar; MediaPipe for in-browser hand-tracking", "Vercel"],
            ["API", "Node.js + Express; ONNX Runtime for the Practice recognition model", "Railway"],
            ["Database", "MySQL (signs, sign_videos, sign_poses, app_settings)", "Railway"],
            ["Demonstration videos", "MP4 files", "Cloudflare R2"],
            ["Sign Chat service", "Python + FastAPI; sign-to-text, text-to-sign and dialogue models; needs a GPU", "Google Colab (Section 4)"],
        ],
        [2300, 5000, 2672],
    )

    d.h1("4. Environments & access")
    d.table(
        ["Environment", "Component", "Address"],
        [
            ["Production", "Client (Vercel)", LIVE + " (public, no login)"],
            ["Production", "API (Railway)", API],
            ["Production", "Database", "MySQL on Railway (shared with the API service)"],
            ["Production", "Sign Chat service", "Google Colab GPU runtime behind a temporary trycloudflare.com URL (below)"],
            ["Preview", "Client", "Vercel builds a preview for every non-master branch (e.g. staging)"],
            ["Local dev", "Client + API", "cd website && npm run install:all && npm run seed && npm run dev (client :5173, API :4000)"],
            ["Local dev", "Sign Chat (no GPU)", "cd sign_chat/sign_chat_backend && SIGNCHAT_MOCK=1 SIGNCHAT_DIALOGUE=echo uvicorn signchat.server:app --port 8000"],
        ],
        [1700, 2300, 5972],
    )
    d.h2("Sign Chat hosting")
    d.para(
        "Sign Chat has no permanent host yet. It runs from sign_chat/sign_chat_backend/colab_server.ipynb on a Colab "
        "A100 (an L4 with High-RAM also works; peak GPU memory is about 15 GB). Run sections 1-4; section 4 has "
        "finished when it prints “models loaded”. Section 3b (a 10-minute precision check) may be skipped, and "
        "section 6 must not be run while the site is in use: it restarts the server, leaving Sign Chat down for about "
        "two minutes."
    )
    d.bullets([
        "The URL is handed over automatically. Each time section 3 opens a tunnel, it reports the new URL to the API (PUT /api/signchat-url, authorised by the SIGNCHAT_UPDATE_TOKEN secret), and the website reads it whenever Sign Chat is opened. Restarting Colab needs no code change and no redeploy. Section 3 prints “website updated: 200” when the hand-off worked.",
        "One-time setup: the same secret value is stored as the Railway variable SIGNCHAT_UPDATE_TOKEN and as a Colab secret of the same name (key icon in Colab's left sidebar, with Notebook access switched on).",
        "Model weights (about 13 GB) are in the Google Drive of ruizisong11@gmail.com, folders auslan_work/ and smplx_models/. The notebook mounts the Drive of the account running it, so either run it from that account or share both folders and add them as shortcuts at the top of My Drive. They are not in the repository -- copy them to storage the sponsor owns at handover.",
        "While the runtime is up, anyone with the URL can call the service. Stop the runtime when it is not needed.",
    ])
    d.figure("signchat.png", "Top: how the website finds the Sign Chat service after a restart. Bottom: one conversation turn.")
    d.h2("Deployment triggers")
    d.para(
        "Both platforms deploy from the master branch automatically: Vercel builds the client, and Railway (project "
        "“pure-kindness”, service at " + API + ") redeploys the API. Pushes to other branches produce Vercel previews only."
    )
    d.figure("release.png", "The release flow: work moves develop → staging → master, and master deploys itself.")

    d.h1("5. Running & operating the product")
    d.para(
        "Client: a static Vite build, auto-deployed by Vercel on every push to the master branch. API: a long-running "
        "Node process on Railway, started with npm start and restarted automatically on crash; GET /api/health returns "
        "{\"status\":\"ok\"} when it is up. The Practice model files (recognition/models, about 14 MB) are committed to "
        "the repository, so nothing needs fetching at deploy time. The 3D avatar (client/public/smplx.glb, 6.2 MB) and "
        "the home page cover clips (client/public/hero) are static files deployed with the client."
    )
    d.h2("Required environment variables")
    d.table(
        ["Variable", "Where", "Purpose"],
        [
            ["PORT", "Server", "Port the Express server listens on (default 4000)."],
            ["DB_HOST / DB_PORT", "Server", "MySQL host and port."],
            ["DB_USER / DB_PASSWORD", "Server", "MySQL credentials -- never commit real values."],
            ["DB_NAME", "Server", "Database name (auslan_learning)."],
            ["CORS_ORIGIN", "Server", "The exact origin allowed to call the API. If unset it defaults to “*” (any origin) -- always set it in production."],
            ["VIDEO_BASE_URL", "Server", "Base URL used to resolve a sign's demonstration video when its row has no explicit video_url."],
            ["RECOGNITION_MODELS", "Server", "Optional path to the recognition model folder; defaults to the repository's recognition/models."],
            ["SIGNCHAT_UPDATE_TOKEN", "Server", "Shared secret the Colab notebook uses to report Sign Chat's URL (same value as the Colab secret). If unset, reporting is switched off (503)."],
            ["VITE_API_BASE_URL", "Client", "Base URL the client calls for the API, including /api (e.g. " + API + "/api)."],
            ["VITE_SIGNCHAT_API_BASE_URL", "Client", "Fallback Sign Chat URL, used only when the API has none on record (website/client/.env.production)."],
        ],
        [2700, 1100, 6172],
    )
    d.para(
        "Secret values live only in each platform's project settings (Vercel / Railway environment variables, Colab "
        "secrets), never in source control; the local .env files are git-ignored and .env.example lists the keys with "
        "no values."
    )

    d.h1("6. Data management")
    d.para(
        "The database has four tables: signs (one row per sign -- gloss, definitions, usage notes, tags, keywords and a "
        "source provenance string, several columns stored as native JSON), sign_videos (one or more demonstration "
        "videos per sign), sign_poses (one gzipped avatar motion clip per sign; a sign without a row simply has no 3D "
        "avatar) and app_settings (small key/value settings -- today only the current Sign Chat URL; created "
        "automatically the first time the notebook reports one)."
    )
    d.h2("Loading data")
    d.bullets([
        "npm run seed -- loads a small set of clearly-labelled placeholder demo signs (local development only).",
        "npm run import:real -- imports the real dataset from server/data/signs.json and server/data/sign_videos.json.",
        "Avatar motion: python scripts/convert_poses.py <npz dir> <out dir> turns fitted SMPL-X .npz files into .json.gz, then node scripts/importPoses.js <out dir> upserts them into sign_poses, matched by gloss (run from website/server; it writes to whichever database the server's .env points at).",
    ])
    d.h2("Backup & restore")
    d.para(
        "Railway backs up the production MySQL volume automatically. To restore, open the MySQL service in the Railway "
        "dashboard, go to its Backups tab, pick a backup and choose Restore; the API reconnects on its own once the "
        "database is back. For a copy outside Railway (e.g. before a risky data import), take a manual dump:"
    )
    d.code("mysqldump -h <DB_HOST> -P <DB_PORT> -u <DB_USER> -p auslan_learning > backup.sql\nmysql -h <DB_HOST> -P <DB_PORT> -u <DB_USER> -p auslan_learning < backup.sql")

    d.h1("7. Security, privacy & standards")
    d.para(
        "There are no end-user accounts. The sign endpoints are read-only; the Practice endpoints score landmark data "
        "without storing it. The only endpoint that writes is PUT /api/signchat-url, which requires the "
        "SIGNCHAT_UPDATE_TOKEN secret (compared in constant time) and accepts only an https URL. Secrets exist only as "
        "platform environment variables and Colab secrets."
    )
    d.figure("practice.png", "Practice keeps the camera picture on the learner's device: only landmark numbers reach the API, and nothing is stored.")
    d.h2("Privacy -- Australian Privacy Principles (Privacy Act 1988)")
    d.table(
        ["Principle", "How HandMirror complies"],
        [
            ["APP 1 -- open and transparent management", "This document and the Product Document describe exactly what each feature sends and keeps. Recommended: a short privacy notice in the site footer."],
            ["APP 3 -- collection", "Only what a feature needs: landmark numbers for Practice, the message for Sign Chat. No names, emails, accounts or tracking cookies."],
            ["APP 5 -- notice of collection", "Practice and Sign Chat run only after the learner starts them. Recommended: one line on Sign Chat saying messages go to the model service."],
            ["APP 6 -- use and disclosure", "Data is used only to produce the score or the reply, and is not shared with anyone else."],
            ["APP 11 -- security of personal information", "HTTPS everywhere; Practice data is not stored; uploaded Sign Chat clips are deleted when the request ends; conversation text is held in memory for the session only."],
            ["APP 12 / 13 -- access and correction", "Nothing is stored about an identifiable learner; saved lists live in the learner's own browser and can be cleared there."],
        ],
        [3300, 6672],
    )
    d.h2("Security -- OWASP Top 10 (2021)")
    d.table(
        ["Risk", "Status in HandMirror"],
        [
            ["A01 Broken access control", "No accounts or private data. The one write endpoint (PUT /api/signchat-url) requires a secret token."],
            ["A02 Cryptographic failures", "All traffic is HTTPS (Vercel, Railway, Cloudflare tunnel). No secrets in the repository."],
            ["A03 Injection", "Every SQL query uses mysql2 parameter placeholders; the reported Sign Chat URL is parsed and reduced to its https origin."],
            ["A04 Insecure design", "Data minimisation by design: Practice sends landmarks, never video."],
            ["A05 Security misconfiguration", "Open items: CORS falls back to “*” when CORS_ORIGIN is unset, and there is no security-headers middleware (helmet). Set CORS_ORIGIN and add helmet before wider use."],
            ["A06 Vulnerable and outdated components", "Run npm audit on client and server regularly; dependencies are pinned in package-lock.json."],
            ["A07 Identification and authentication failures", "Not applicable (no logins). Rotate SIGNCHAT_UPDATE_TOKEN if it is ever exposed."],
            ["A09 Logging and monitoring", "Railway and Vercel keep request and error logs; GET /api/health supports uptime monitoring."],
            ["A10 Server-side request forgery", "The API never fetches a URL supplied by a user; the Sign Chat URL is only stored and handed to browsers."],
        ],
        [3300, 6672],
    )
    d.h2("Accessibility -- WCAG 2.1 Level AA")
    d.para("Every page was checked with axe-core against WCAG 2.0-2.2 A/AA rules (Maintenance Document, Section 10). Keep new pages at the same level.")
    d.h2("Known limitations")
    d.bullets([
        "No rate-limiting middleware: each POST /api/recognize call runs model inference on the API server's CPU, so a flood of requests could slow the whole API. Add express-rate-limit (tighter on /api/recognize).",
        "Request bodies up to 5 MB are accepted (a Practice capture is about 190 KB).",
        "While Colab is running, the Sign Chat tunnel URL is public and unauthenticated.",
        "The Practice score threshold was calibrated on studio (green-screen) footage; home webcams, lighting and framing can lower scores.",
        "Sign Chat's sign-to-text and text-to-sign models have not been validated by fluent signers; they do best on short, everyday sentences.",
    ])

    d.h1("8. API reference summary")
    d.table(
        ["Endpoint", "Notes"],
        [
            ["GET /api/health", "Liveness check; returns {\"status\":\"ok\"}."],
            ["GET /api/signs", "Search/list signs. Optional: query (search across gloss, tags, keywords, definitions), tag, level (beginner / intermediate / advanced), sort, page, pageSize (default 6, max 50)."],
            ["GET /api/signs/tags", "Topic categories with counts. Optional exclude (comma-separated sign ids) so counts reflect what is still left to learn."],
            ["GET /api/signs/sample", "Random signs from one topic for the personalise wizard: tag, count (max 50), exclude."],
            ["GET /api/signs/:id", "Single sign with its demonstration videos and estimated level. 404 if the id doesn't exist."],
            ["GET /api/signs/:id/pose", "The sign's avatar motion clip (SMPL-X joint rotations per frame). 404 if the sign has none."],
            ["POST /api/recognize/verify", "Practice scoring. Body {word, capture}; returns matched, a 0–100 score, a tier (excellent / close / needs_practice) and a feedback sentence."],
            ["POST /api/recognize/identify", "Sign-to-search. Body {capture}; returns the four closest signs with a calibrated confidence."],
            ["GET /api/recognize/vocabulary", "Which signs the recognition model knows (?detail=1 adds the Signbank cross-reference)."],
            ["GET /api/signchat-url", "The Sign Chat service's current URL and when it was reported; 404 if none yet."],
            ["PUT /api/signchat-url", "Header Authorization: Bearer <SIGNCHAT_UPDATE_TOKEN>, body {url}. 401 wrong token, 400 not an https URL, 503 token not configured."],
            ["Sign Chat service", "POST /api/chat/text, POST /api/chat/sign (multipart video upload), DELETE /api/session/:id -- see sign_chat/sign_chat_backend/README.md."],
        ],
        [3000, 6972],
    )
    d.para("The landmark “capture” format and error codes (400 unusable_capture, 404 unknown_word, 503 model_unavailable) are specified in recognition/API.md.")

    d.h1("9. Troubleshooting guide")
    d.table(
        ["Symptom", "Likely cause / fix"],
        [
            ["Server won't start", "Check DB_HOST/DB_PORT/DB_USER/DB_PASSWORD/DB_NAME point to a reachable MySQL instance."],
            ["CORS error in the browser console", "CORS_ORIGIN on the server doesn't match the client's actual deployed origin."],
            ["Videos don't play / show black", "VIDEO_BASE_URL isn't set, or the Cloudflare R2 bucket it points to is unreachable."],
            ["Search returns zero results", "The database hasn't been seeded or imported yet -- run npm run seed or npm run import:real."],
            ["Practice says “isn't switched on for this site yet”", "The API returned 503 model_unavailable: recognition/models is missing from the deployment, or RECOGNITION_MODELS points at the wrong folder."],
            ["Practice says “Didn't catch enough movement”", "400 unusable_capture: both hands weren't in frame, or the capture was too short. Ask the learner to step back so both shoulders are visible."],
            ["“There's no 3D avatar for this sign yet”", "The sign has no sign_poses row -- import its motion clip (Section 6)."],
            ["3D avatar panel stays blank", "smplx.glb failed to load, or the browser has WebGL disabled; check the browser console."],
            ["Sign Chat shows an error instead of a reply", "The Colab runtime is not running, or the models are still loading (about two minutes after a restart). Start the notebook (sections 1-4), then reload the Sign Chat page."],
            ["Section 3 prints “website NOT updated: Secret SIGNCHAT_UPDATE_TOKEN does not exist”", "Add the Colab secret (key icon, Notebook access on), then re-run the report line or section 3."],
            ["Section 3 prints 503 “SIGNCHAT_UPDATE_TOKEN is not set on the server” / 401 “Bad token”", "Set the Railway variable, or make the Railway and Colab values identical."],
            ["Section 3 shows a URL of https://api.trycloudflare.com", "Cloudflare refused the tunnel (usually rate limiting). The notebook retries automatically; if all attempts fail, wait 10-15 minutes and run section 3 again."],
            ["Sign Chat repeats your text back and the avatar barely moves", "The client is talking to a mock server (SIGNCHAT_MOCK=1, echo dialogue), not the real models."],
        ],
        [3300, 6672],
    )

    d.h1("10. Knowledge transfer plan")
    d.para(
        "The handover is designed so a sponsor's staff can run and change HandMirror without the original team. It "
        "happens in four steps, each led by the team members below and backed by these documents."
    )
    d.table(
        ["Step", "Audience", "Content", "Material"],
        [
            ["1. Product walkthrough (1 hour)", "Sponsor's product owner, support staff", "Live demo of every feature in Section 2; who the product serves; licence and privacy position", "Product Document; this document, Section 2"],
            ["2. Operations session (1.5 hours)", "Support staff", "Deploying, environment variables, starting Sign Chat on Colab, backups and restore, the troubleshooting table", "This document, Sections 3-9"],
            ["3. Code walkthrough (2 hours)", "Technical team", "Codebase map, the recognition and avatar pipelines, the release flow, how to add signs and journeys", "Maintenance Document; repository READMEs"],
            ["4. Account transfer", "Sponsor's administrator", "Ownership or admin rights moved to the sponsor (list below); secrets rotated afterwards", "Checklist below"],
        ],
        [2200, 2000, 3600, 2172],
    )
    d.para("Accounts and assets to transfer:")
    d.bullets([
        "GitHub repository " + REPO + " (transfer ownership or add the sponsor as admin).",
        "Vercel project “handmirror” (team te-38) and Railway project “pure-kindness” (API and MySQL).",
        "Cloudflare R2 bucket with the demonstration videos.",
        "Google Drive folders with the Sign Chat model weights (auslan_work/, smplx_models/) and the Colab notebook.",
        "The Project Governance Portfolio on Google Drive (" + GOVERNANCE + ").",
        "After transfer: rotate DB_PASSWORD and SIGNCHAT_UPDATE_TOKEN so the team's copies stop working.",
    ])

    d.h1("11. Support contacts")
    d.table(
        ["Name", "Role", "Contact"],
        [
            ["Dingao Wang", "Lead developer -- first contact for any issue", "wangdingao5337@gmail.com"],
            ["Ruizi Song", "Holds the Sign Chat model weights (Google Drive)", "ruizisong11@gmail.com"],
        ],
        [2400, 4300, 3272],
    )
    d.closing("HandMirror — Support Document · Team TE38 (BitDance) · FIT5120")


def maintenance(d: Doc):
    d.header("HANDOVER PACKAGE — MAINTENANCE DOCUMENT")
    d.h1("1. Purpose")
    d.para(
        "This document is for a future technical team extending HandMirror -- adding features, fixing bugs, retraining "
        "models, or building out the epics still on the roadmap. It assumes the reader is a competent developer new to "
        "this specific codebase."
    )
    d.h1("2. Codebase map")
    d.table(
        ["Path", "What's there"],
        [
            ["website/client/src/pages", "One component per route: Home, SignLibrary, SignDetail, Scenario list and journey (ScenarioCategoryPage), ToLearn, LearnedSigns (shown as “Learnt”; the route stays /learned), Conversation (Sign Chat)."],
            ["website/client/src/components", "Shared UI. Key pieces: SmplxAvatar (3D avatar), SignDemonstration (video / 3D avatar tabs), PracticeVerify and SignCaptureModal (webcam capture), GestureSearch (sign-to-search), PersonalizeSessionPanel + PieRatioChart (personalise wizard), SavedSignsPage + TimeGroupedList (To Learn / Learnt), ResultCard (sign card with the + / ✓ buttons), BulkActions (the Select signs bar), WordPreviewModal (scenario hover preview)."],
            ["website/client/src/api", "Typed fetch wrappers: signs.ts (dictionary), recognize.ts (Practice / sign-to-search), signChat.ts (Sign Chat service; asks the API for the service's current URL first)."],
            ["website/client/src/hooks", "useSelection (multi-select state), useHoverPreview, useSignSearch, useTags and other small hooks."],
            ["website/client/src/lib", "landmarks.ts (MediaPipe capture with motion-based auto-stop), smplx.ts (avatar pose format), scenarioCategories.ts (the 34 journeys and their words), pictograms.ts (journey icons), learnedSigns / toLearnSigns / idListStore (localStorage lists) and signLists.ts (the rule that a sign is on one list at most)."],
            ["website/client/src/assets/pictograms", "The 36 journey pictograms (Human Pictogram 2.0, recoloured)."],
            ["website/client/public/smplx.glb", "The 3D avatar: SMPL-X body with texture, hair and a “smile” morph target."],
            ["website/client/src/styles/theme.css", "The whole design system in one file -- colour/type/radius tokens plus every page's styles."],
            ["website/server/src/routes", "signs.js (dictionary and avatar poses), recognize.js (Practice and sign-to-search), signchat.js (stores and serves the Sign Chat URL)."],
            ["website/server/src/recognition", "features.js (landmarks → model input, a port of the Python training pipeline) and model.js (ONNX encoder and template bank)."],
            ["website/server/src/utils", "formatSign / pagination / video helpers, difficulty.js (level heuristic), popularity.js (popularity sort)."],
            ["website/server/scripts", "importRealData.js, convert_poses.py and importPoses.js (data loading -- see the Support Document)."],
            ["recognition/", "Model files, the interface specification (API.md), deployment notes (ASSETS.md), build scripts and the parity fixtures (test/golden.json)."],
            ["sign_chat/sign_chat_backend", "The Sign Chat model service (FastAPI), its README, Colab launcher and tests."],
        ],
        [3100, 6872],
    )
    d.figure("arch.png", "The deployed system. Section 8 shows how a change reaches production.")

    d.h1("3. Adding or editing content")
    d.h2("Signs")
    d.para("gloss has a UNIQUE constraint, so an insert can double as an upsert:")
    d.code(
        "INSERT INTO signs (gloss, definitions, usage_notes, tags, keywords, source)\n"
        "VALUES ('THANK YOU',\n"
        "  CAST('[{\"partOfSpeech\":\"Interjection\",\"senses\":[\"A polite expression of gratitude.\"]}]' AS JSON),\n"
        "  CAST('[\"Touch fingertips to chin.\", \"Move hand forward toward the other person.\"]' AS JSON),\n"
        "  CAST('[\"greeting\"]' AS JSON),\n"
        "  CAST('[\"thanks\",\"cheers\"]' AS JSON),\n"
        "  'Auslan Signbank (pending verification)');"
    )
    d.para(
        "tags are visible classification chips (a sign can have several); keywords are search-only synonyms that are "
        "never shown in the UI. When in doubt: a genuine topic/category is a tag, another word someone might type to "
        "find the sign is a keyword."
    )
    d.h2("Scenario journeys")
    d.para(
        "Journeys and their stops live in client/src/lib/scenarioCategories.ts (words are matched to signs by gloss). "
        "A new journey also needs an entry in client/src/lib/pictograms.ts mapping its id to a pictogram file."
    )
    d.h2("Avatar motion")
    d.para("Add or replace a sign's 3D motion with convert_poses.py and importPoses.js (Support Document, Section 6). The clip must be fitted to the same SMPL-X body the avatar uses (Section 6).")

    d.h1("4. Extending the API")
    d.bullets([
        "One express.Router per resource, mounted in src/index.js -- follow the existing pattern in src/routes/signs.js for a new resource.",
        "Always use parameterized ? placeholders via mysql2 -- never string-concatenate user input into SQL.",
        "Map DB rows to the API's camelCase JSON shape through a small helper in src/utils/ (see formatSign.js), keeping snake_case DB columns out of the public response shape.",
    ])

    d.h1("5. Practice recognition")
    d.para(
        "MediaPipe Holistic runs in the browser and records 48 landmarks per frame (arms and both hands); the capture "
        "stops itself shortly after the hands come to rest. The API turns the landmarks into a fixed 96 × 98 tensor "
        "(features.js), runs the encoder (encoder.onnx, 256-dimensional output) and compares the result with the "
        "target sign's 12 stored templates by cosine distance."
    )
    d.figure("practice.png", "The Practice pipeline, split between the browser and the API.")
    d.bullets([
        "Verify (Practice): matched when the distance is below τ = 0.529 -- the equal-error point on the validation split (EER 0.73%). The 0–100 score is a rescaling of that distance anchored so 60 is exactly the pass line.",
        "Left- or right-handed: the attempt is scored as captured and mirrored (other hand leading), and the closer result is kept. This slightly raises the chance of accepting a wrong sign.",
        "Identify (sign-to-search): distance to every sign; the four closest are returned with a confidence calibrated on the validation split.",
        "Any change to features.js must keep npm run test:recognition (in website/server) passing -- it checks the JavaScript pipeline against fixtures exported from the Python original.",
        "After retraining, rebuild the template bank with recognition/scripts/buildBank.mjs and follow recognition/ASSETS.md. The full contract, including capture guidance for the UI, is recognition/API.md.",
    ])

    d.h1("6. The 3D avatar")
    d.para(
        "SmplxAvatar loads /smplx.glb: the SMPL-X neutral body with one bone per SMPL-X joint, named after the joint, "
        "with identity rest rotations. Each frame's SMPL-X axis-angle rotations are applied straight onto those bones -- "
        "no retargeting -- so the avatar reproduces the fitted motion, including where the hands touch. On top of the "
        "motion the component adds eyes that follow the camera within ±8°, a resting smile on Sign Chat only (a morph "
        "target), and an idle pose set to the median trunk and head angle of the signing clips."
    )
    d.para(
        "Pose clips are JSON: {fps, frames, smplx: {global_orient, body_pose (21 joints × 3), left_hand_pose and "
        "right_hand_pose (15 × 3 each)}}, stored gzipped in sign_poses. They come from SMPL-X fits of the dataset's "
        "signers; motion fitted to any other body would not line up."
    )
    d.para(
        "The scripts that fit SMPL-X to the dataset and build smplx.glb (body, textures, clothes, hair) are in the "
        "Project Governance Portfolio as HandMirror-avatar-build-scripts.zip, with a README listing the order to run "
        "them and the input files. The SMPL-X model, BEDLAM skin and Meshcapade textures are not in the zip, because "
        "their licences forbid redistribution; download them from their official sites (Product Document, Section 4)."
    )

    d.h1("7. Sign Chat service")
    d.para(
        "A separate Python service (sign_chat/sign_chat_backend). A typed message goes to a dialogue model; a signed "
        "message is first turned into text by a sign-to-text model; the reply text then goes through a text-to-sign "
        "model that produces the avatar motion the client plays. Its README covers the models, the API, a mock mode for "
        "front-end work without a GPU (SIGNCHAT_MOCK=1), and the tests (python -m pytest -q tests)."
    )
    d.para(
        "Because Colab gives the service a new URL on every restart, the URL is looked up at run time: notebook "
        "section 3 reports it to the API (PUT /api/signchat-url, routes/signchat.js, stored in app_settings), and "
        "client/src/api/signChat.ts reads it once per page load (GET), falling back to VITE_SIGNCHAT_API_BASE_URL. "
        "A permanent GPU host with a fixed address would make the hand-off unnecessary; it can stay as a fallback."
    )
    d.figure("signchat.png", "Finding the service, then one conversation turn.")

    d.h1("8. Release process")
    d.para(
        "Work is committed on develop, fast-forwarded to staging (Vercel builds a preview to check), then merged into "
        "master, which deploys the client (Vercel) and the API (Railway) automatically. Before merging to master: "
        "npm run build and npm run lint in website/client, npm run test:recognition in website/server, and a quick "
        "check of the preview. Never push straight to master."
    )
    d.figure("release.png", "Branches and what each one deploys.")

    d.h1("9. Coding conventions observed")
    d.bullets([
        "Client: TypeScript in strict mode (React 19 + Vite); lint with npm run lint (oxlint).",
        "Server: plain CommonJS Node.js (no TypeScript, no ORM -- raw parameterized SQL via mysql2).",
        "camelCase at the API/JSON boundary; snake_case for DB columns -- translated by the format helpers, never leaked across that boundary in either direction.",
        "Design system lives entirely in client/src/styles/theme.css: IBM Plex Mono for display/labels, IBM Plex Sans for body text, with colour/spacing/radius as CSS custom properties.",
        "Comments explain why a non-obvious choice was made; keep that habit when changing the recognition or avatar code, where the reasons are not visible in the code itself.",
        "Interface wording uses Australian English (“Learnt”, “practise” as the verb); code identifiers keep their original names (learnedSigns, /learned) so saved lists and links keep working.",
    ])

    d.h1("10. Standards & compliance")
    d.h2("Data and asset licensing")
    d.para(
        "Every sign record's source field must stay accurate. The licences of all third-party assets are listed in the "
        "Product Document, Section 4. The restricted ones are Creative Commons 4.0 non-commercial (or equivalent): the "
        "product may be used for non-commercial purposes only, and any commercial use needs a separate licence from "
        "each rights holder first. The SMPL-X model may not be redistributed, so it is not in the repository's build scripts."
    )
    d.h2("Privacy and security")
    d.para(
        "Changes must keep the product within the Australian Privacy Principles (Privacy Act 1988) and the OWASP Top 10 "
        "measures mapped in the Support Document, Section 7 -- in particular: no new collection of personal "
        "information without a notice, parameterised SQL only, and authentication on any new write endpoint."
    )
    d.h2("Accessibility")
    d.para(
        "Conformance level: WCAG 2.1 Level AA. Checked on 13 October 2026 with axe-core 4.10.2 (all WCAG 2.0, 2.1 and "
        "2.2 A/AA rules) in Chrome on every page: home, library, sign detail, scenarios, scenario categories, To Learn, "
        "Learnt and Sign Chat. The demonstration videos have no audio, so captions are not required."
    )
    d.para(
        "Keep it at AA when changing the interface: text needs a contrast of at least 4.5:1 (3:1 for large text), "
        "every control and progress bar needs an accessible name, and re-run axe on any page you change."
    )
    d.h2("Security hygiene going forward")
    d.bullets([
        "Secrets stay in environment variables only -- never hard-coded or committed.",
        "Add rate limiting (stricter on /api/recognize) and helmet before wider use; set CORS_ORIGIN explicitly.",
        "Any future write-capable endpoint must add real authentication/authorisation.",
        "Run npm audit periodically on both client and server and address high-severity advisories.",
    ])

    d.h1("11. Roadmap / open epics")
    d.table(
        ["Item", "Status"],
        [
            ["Basic quiz / self-check (Epic 3)", "Not yet built."],
            ["Family mode for parent-child signing (Epic 6)", "Not yet built."],
            ["Permanent hosting for the Sign Chat service", "Runs from a Colab notebook today, with weights in a team member's Google Drive (Support Document, Section 4). Needs a GPU host with at least 24 GB (e.g. an L4) and a fixed URL."],
            ["Facial expression on the avatar", "Auslan grammar uses the face (e.g. raised brows for questions); the avatar's face is not driven yet."],
            ["Validation by fluent signers", "Practice thresholds and Sign Chat output have not yet been checked by Auslan signers."],
        ],
        [4300, 5672],
    )

    d.h1("12. Known technical debt")
    d.bullets([
        "ScenarioLearningPage (route /scenarios/:id) and its story data (scenarioStories.ts, useScenarioProgress, ScenarioWordLink) are no longer linked from anywhere and can be removed.",
        "Unused leftovers: the @pixiv/three-vrm dependency, client/public/avatar.vrm (untracked), and style rules in theme.css for the previous home page and scenario sidebar.",
        "smplx.glb is 6.2 MB, mostly the hair mesh; simplifying the hair would speed up the first avatar load.",
        "The home page plays up to 33 short videos at once (the cover plus the sign strip); consider fewer clips or playing them only on hover for low-end devices.",
        "The four cover clips in client/public/hero are the original 1080p Pexels downloads (2.7-9 MB each, 22 MB in total; one plays at a time). Re-encoding them to 720p with ffmpeg would cut them to about a quarter of that.",
        "Automated tests: the recognition parity test (server) and the Sign Chat service's pytest suite. There are no client tests and no CI pipeline beyond Vercel/Railway building on push.",
    ])
    d.closing("HandMirror — Maintenance Document · Team TE38 (BitDance) · FIT5120")


ITERATION_ITEMS = ["Presentation", "Iteration Analysis & Design Report", "Data Management Plan", "Agenda, Minutes and Reflection",
                   "Usability Video", "Retrospectives", "Iteration project supporting documents"]


def checklist(d: Doc):
    d.header("HANDOVER PACKAGE", "Handover Checklist and Sign-Off")
    d.para(
        "We, Team TE38 (“BitDance”), acknowledge that we have provided all the handover documents listed below to the "
        "studio mentors named below, via the Project Governance Portfolio (Google Drive link below)."
    )
    d.table(
        ["Item", "Details"],
        [
            ["Team members", ", ".join(m[0] for m in TEAM)],
            ["Studio mentors", MENTORS],
            ["Project Governance Portfolio", GOVERNANCE],
        ],
        [3000, 6972],
    )
    d.h1("Project Governance Portfolio -- folder contents")
    d.para("Every item below is in the Project Governance Portfolio on Google Drive (confirmed by the team, 15 October 2026).")
    ok = "✓ In the Portfolio"
    d.h2("Monash University Use")
    d.table(
        ["Item", "Status / location"],
        [
            ["Project Proposal", ok + " (Inclusive_Auslan_Learning_Assistant_Proposal.pptx)"],
            ["Lotus Blossom", ok],
            ["Industry Mentor Pitch", ok],
            ["Customer Journey Map", ok],
            ["Empathy Map", ok],
            ["Persona", ok + " (primary persona “Sarah Thompson”, Product Document Section 1)"],
            ["Prototype", ok],
            ["Sitemap", ok],
            ["Ethics Canvas", ok + " (mapped to the ACS Code of Conduct in the Product Document, Section 4)"],
            ["System Architecture Diagrams", ok + " (the architecture diagram is also in the Support and Maintenance Documents)"],
            ["Other project supporting documents", ok],
        ],
        [4000, 5972],
    )
    d.h2("Iterations folder")
    for n in (1, 2, 3):
        d.table(["Iteration " + str(n), "Status"], [[item, ok] for item in ITERATION_ITEMS], [6000, 3972])
    d.table(["Other documents", "Status"], [["Other iteration documents", ok]], [6000, 3972])
    d.h2("Code and Database folder")
    d.table(
        ["Item", "Status / location"],
        [
            ["Data pre-processing", ok + "; scripts in the repository: website/server/scripts (importRealData.js, convert_poses.py, importPoses.js), recognition/scripts, and HandMirror-avatar-build-scripts.zip"],
            ["SQL script for database", ok + " (website/server/schema.sql)"],
            ["Final Data Management Plan", ok],
        ],
        [4000, 5972],
    )
    d.h2("Project Support folder")
    d.table(["Item", "Status"], [[i, ok] for i in ("Product video", "Stand-up meetings", "Expo content", "Team info doc")], [6000, 3972])
    d.h2("Security folder")
    d.table(["Item", "Status"], [[i, ok] for i in ("Security Plan", "Vulnerability Assessment Report")], [6000, 3972])
    d.h2("Handover Documents")
    d.table(
        ["Item", "Status / location"],
        [
            ["Documentation -- Product Document", ok + " (01-Product-Document.docx)"],
            ["Documentation -- Support Document", ok + " (02-Support-Document.docx)"],
            ["Documentation -- Maintenance Document", ok + " (03-Maintenance-Document.docx)"],
            ["Code -- backup", "GitHub repository: " + REPO + "; avatar build scripts: HandMirror-avatar-build-scripts.zip"],
        ],
        [4000, 5972],
    )
    d.h1("Sign-off")
    d.para("Team TE38 (“BitDance”)")
    d.table(["Name", "Signature", "Date"], [[m[0], "", ""] for m in TEAM], [3800, 3800, 2372])
    d.h2("Monash mentor acceptance")
    d.table(["Name", "Signature", "Date"], [["Afsana", "", ""], ["Ted", "", ""]], [3800, 3800, 2372])
    d.closing("HandMirror — Handover Checklist and Sign-Off · Team TE38 (BitDance) · FIT5120")


if __name__ == "__main__":
    root = Path(sys.argv[1])
    for name, build in [
        ("01-Product-Document.docx", product),
        ("02-Support-Document.docx", support),
        ("03-Maintenance-Document.docx", maintenance),
        ("04-Handover-Checklist-and-Sign-off.docx", checklist),
    ]:
        doc = Doc(root / name)
        build(doc)
        doc.save()
        print("wrote", name)
